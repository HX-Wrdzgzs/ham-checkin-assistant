"""Excel 历史批量导入：单文件/多文件/整个文件夹。

去重：file_hash + sheet_name + row_number（raw_imports UNIQUE 约束）+ record_hash。
同一文件重复导入不会产生重复记录。
"""
from __future__ import annotations

import json
import re
from datetime import datetime
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook

from database.models import Checkin, RawImport
from database.repository import Repository
from excel import template
from providers.base import DataProvider


class ExcelImportProvider(DataProvider):
    name = "excel_import"

    def __init__(self, repo: Repository) -> None:
        self.repo = repo

    def check_updates(self) -> dict:
        return {"has_new": False}

    def fetch_records(self) -> list[dict]:
        return []

    def normalize_record(self, raw: dict) -> dict:
        return raw

    # ---------- 导入 ----------
    def import_file(self, path: Path, standardizer=None) -> tuple[int, int, str]:
        """返回 (导入数, 跳过数, 说明)。原子导入：失败整体回滚（任务书第三阶段 #1/#2）。"""
        path = Path(path)
        if not path.exists():
            return 0, 0, "文件不存在"
        if path.suffix.lower() == ".xls":
            return 0, 0, "暂不支持 .xls，请另存为 .xlsx 后导入（任务书第三阶段 #4）"
        if path.suffix.lower() != ".xlsx":
            return 0, 0, "仅支持 .xlsx 文件"
        fh = self.repo.file_hash(str(path))
        if self.repo.import_job_completed("excel_import", fh):
            return 0, 0, "该文件已导入过（已跳过）"

        # 1) 解析（流式，不写库；任何异常都不产生半导入）
        try:
            wb = load_workbook(path, data_only=True, read_only=True)
        except Exception as e:  # noqa: BLE001
            self.repo.mark_import_job("excel_import", fh, "failed", error=str(e))
            return 0, 0, f"文件无法读取：{e}"
        try:
            prepared: list[dict] = []
            first_date = ""
            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                rows_iter = ws.iter_rows(values_only=True)
                # 缓冲前 12 行识别表头（find_header_row 只查前 10 行）
                head = []
                for _ in range(12):
                    try:
                        head.append(next(rows_iter))
                    except StopIteration:
                        break
                header_idx, mapping = template.find_header_row(head)
                if header_idx is None or not mapping.get("callsign"):
                    continue
                # 表头之后的行 = 缓冲中剩余 + 流式剩余（真正 streaming，不整体加载）
                for ri, row_vals in enumerate(head[header_idx + 1:], start=header_idx + 2):
                    rec = self._row_to_record(mapping, row_vals, fh, sheet_name, ri)
                    if rec is None:
                        continue
                    if not first_date:
                        first_date = rec.get("date", "")
                    prepared.append(rec)
                ri = len(head)
                for row_vals in rows_iter:
                    ri += 1
                    rec = self._row_to_record(mapping, row_vals, fh, sheet_name, ri)
                    if rec is None:
                        continue
                    if not first_date:
                        first_date = rec.get("date", "")
                    prepared.append(rec)
        finally:
            wb.close()
        if not prepared:
            self.repo.mark_import_job("excel_import", fh, "failed",
                                      error="未识别到有效数据行")
            return 0, 0, "未识别到有效数据行"

        date = self._guess_date(path, [p["rec"] for p in prepared], first_date)
        # P0-3：先定有效日期/时间 → 规范化 canonical datetime → 再生成指纹（跨日期不再误去重）
        self._finalize_records(prepared, date)
        raw_imports = [p["raw"] for p in prepared]
        checkins = [self._build_checkin(0, p["rec"], standardizer) for p in prepared]
        try:
            # P1-8：session 创建/结束 + raw_imports + checkins 单事务（失败无 ghost session/半导入）
            imported, skipped = self.repo.import_file_atomic(
                "excel_import", path.stem, date, raw_imports, checkins)
        except Exception as e:  # noqa: BLE001
            # 原子失败：不留 raw_imports/checkins/active session 半导入，也不标记 completed
            self.repo.mark_import_job("excel_import", fh, "failed", error=str(e))
            return 0, 0, f"导入失败（已回滚）：{e}"
        # 重建投影（checkins 为唯一事实源）
        self.repo.rebuild_all_stations()
        self.repo.rebuild_all_profiles()
        self.repo.mark_import_job("excel_import", fh, "completed", imported_count=imported)
        return imported, skipped, f"导入 {imported} 条，跳过 {skipped} 条"

    def import_files(self, paths: Iterable[Path], standardizer=None) -> dict[str, tuple[int, int, str]]:
        out = {}
        for p in paths:
            try:
                out[str(p)] = self.import_file(Path(p), standardizer)
            except Exception as e:  # noqa: BLE001
                out[str(p)] = (0, 0, f"失败：{e}")
        return out

    def import_folder(self, folder: Path, standardizer=None) -> dict[str, tuple[int, int, str]]:
        # 任务书第三阶段 #4：只支持 .xlsx
        files = sorted(folder.glob("*.xlsx"))
        return self.import_files(files, standardizer)

    # ---------- 内部 ----------
    def _finalize_records(self, prepared: list[dict], session_date: str) -> None:
        """确定每行有效日期 → 规范化 canonical datetime → 生成业务指纹（P0-3）。

        指纹 = source + canonical datetime + callsign + qth/device/antenna/power/signal。
        行有日期列用行日期，否则用文件/推断日期（确保不同日期的同名记录不误判重复）。
        """
        from core.datetime_util import normalize_checkin_time

        for p in prepared:
            rec = p["rec"]
            effective_date = rec.get("date") or session_date
            canonical = normalize_checkin_time(effective_date, rec.get("time", ""))
            rec["date"] = effective_date
            rec["canonical_datetime"] = canonical
            rec["record_hash"] = self.repo.record_hash(
                "excel_import", canonical, (rec.get("callsign") or "").upper(),
                rec.get("qth", ""), rec.get("device", ""), rec.get("antenna", ""),
                rec.get("power", ""), rec.get("signal", ""))
            p["raw"].record_hash = rec["record_hash"]

    def _file_imported(self, fh: str) -> bool:
        return self.repo.import_job_completed("excel_import", fh)

    def _row_to_record(self, mapping: dict, row_vals: list, fh: str, sheet: str, ri: int) -> dict | None:
        """解析一行 → {rec, raw}，不写库（原子导入时统一写）。"""
        def get(field: str) -> str:
            col = mapping.get(field)
            if col is None or col >= len(row_vals):
                return ""
            v = row_vals[col]
            return str(v).strip() if v is not None else ""

        callsign = get("callsign")
        if not callsign:
            return None
        date, time = get("date"), get("time")
        # 业务指纹在 _finalize_records 中（先定日期/规范化时间后）生成（P0-3）
        rec = {
            "callsign": callsign,
            "time": time,
            "qth": get("qth"),
            "device": get("device"),
            "antenna": get("antenna"),
            "power": get("power"),
            "signal": get("signal"),
            "date": date,
            "raw_json": json.dumps({k: get(k) for k in
                                    ("sequence", "time", "callsign", "qth", "device", "antenna", "power", "signal")},
                                   ensure_ascii=False),
            "record_hash": "",
        }
        raw = RawImport(
            source="excel_import", source_file=fh, sheet_name=sheet,
            row_number=ri, raw_json=rec["raw_json"], record_hash="",
        )
        return {"rec": rec, "raw": raw}

    @staticmethod
    def _guess_date(path: Path, rows: list[dict], first_date: str) -> str:
        if first_date:
            return first_date
        m = re.search(r"(\d{4}[-_.]?\d{2}[-_.]?\d{2})", path.stem)
        if m:
            return m.group(1).replace("_", "-").replace(".", "-")
        return datetime.fromtimestamp(path.stat().st_mtime).strftime("%Y-%m-%d")

    def _build_checkin(self, session_id: int, rec: dict, standardizer) -> Checkin:
        from core.datetime_util import normalize_checkin_time

        c = Checkin(
            session_id=session_id,
            # 任务书第三阶段 #6 + P0-3：统一 YYYY-MM-DDTHH:MM:SS（含文件/推断日期）
            checkin_time=rec.get("canonical_datetime")
            or normalize_checkin_time(rec.get("date", ""), rec.get("time", "")),
            callsign=(rec.get("callsign") or "").upper(),
            qth_raw=rec.get("qth", ""), device_raw=rec.get("device", ""),
            antenna_raw=rec.get("antenna", ""), power_raw=rec.get("power", ""),
            signal=rec.get("signal", ""),
            source="excel_import",
            raw_input=rec.get("raw_json", ""),
            source_record_id=rec.get("record_hash", ""),
            source_url=str(rec.get("source_url", "")),
        )
        if standardizer:
            c.qth_standard = standardizer.standardize("qth", c.qth_raw) or c.qth_raw
            c.device_standard = standardizer.standardize("device", c.device_raw) or c.device_raw
            c.antenna_standard = standardizer.standardize("antenna", c.antenna_raw) or c.antenna_raw
            c.power_standard = standardizer.standardize("power", c.power_raw) or c.power_raw
        else:
            c.qth_standard = c.qth_raw
            c.device_standard = c.device_raw
            c.antenna_standard = c.antenna_raw
            c.power_standard = c.power_raw
        return c
