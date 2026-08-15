"""回归测试：P1-12 —— 公式注入防护统一（COM 与 openpyxl 共用 excel_safe_text）。"""
from __future__ import annotations

import unittest

from excel.template import excel_safe_text

CASES = ["=1+1", "+CMD", "-1+2", "@SUM(A1:A2)", "=HYPERLINK(...)"]
SAFE = ["BA4XXX", "南京栖霞", "K6", "5W", "59"]


class TestExcelSafeText(unittest.TestCase):
    def test_formula_prefixes_escaped(self):
        for v in CASES:
            self.assertTrue(str(excel_safe_text(v)).startswith("'"),
                            f"{v!r} 必须加单引号按文本处理")

    def test_normal_values_unchanged(self):
        for v in SAFE:
            self.assertEqual(excel_safe_text(v), v)

    def test_non_string_unchanged(self):
        self.assertEqual(excel_safe_text(123), 123)
        self.assertEqual(excel_safe_text(None), None)

    def test_export_uses_excel_safe_text(self):
        """导出 xlsx 中公式前缀被转义（openpyxl 路径）。"""
        import tempfile
        from pathlib import Path

        from openpyxl import load_workbook

        from database.db import connect
        from database.repository import Repository
        from database.seed import seed_default_aliases
        from excel.exporter import export_session
        from database.models import Checkin

        tmp = Path(tempfile.mkdtemp())
        conn = connect(tmp / "e.db")
        repo = Repository(conn)
        seed_default_aliases(repo)
        s = repo.create_session("导出", "2026-08-08")
        repo.add_checkin(Checkin(
            session_id=s.id, sequence_no=1, checkin_time="2026-08-08T20:00:00",
            callsign="=1+1", qth_standard="+南京", device_standard="K6",
            antenna_standard="-1", power_standard="@SUM(1,1)", source="local"))
        try:
            dest = tmp / "export.xlsx"
            export_session(repo, s.id, str(dest))
            wb = load_workbook(dest)
            ws = wb.active
            row = [c.value for c in ws[5]]
            self.assertTrue(str(row[2]).startswith("'"), f"呼号被转义: {row[2]!r}")
            self.assertTrue(str(row[3]).startswith("'"))
            self.assertTrue(str(row[5]).startswith("'"))
        finally:
            conn.close()


if __name__ == "__main__":
    unittest.main()
