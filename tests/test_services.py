"""服务层测试：Excel 导入去重 / 提交撤销 / 365dt 容错 / 导出重建（规格第 100~102 节）。"""
from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from database.db import connect
from database.models import Checkin
from database.repository import Repository
from database.seed import seed_default_aliases
from excel.exporter import export_session
from normalizers.dictionaries import AliasStore
from normalizers.region_index import RegionIndex
from providers.excel_import import ExcelImportProvider
from services.standardizer import Standardizer
from services.sync_service import SyncService


def make_env():
    tmp = Path(tempfile.mkdtemp())
    conn = connect(tmp / "test.db")
    repo = Repository(conn)
    seed_default_aliases(repo)
    store = AliasStore(repo)
    region = RegionIndex("江苏")
    std = Standardizer(store, region)
    return tmp, conn, repo, std


def make_xlsx(path: Path, rows: list[list]):
    wb = Workbook()
    ws = wb.active
    ws.append(["序号", "时间", "呼号", "QTH", "设备", "天线", "功率"])
    for r in rows:
        ws.append(r)
    wb.save(path)


class TestExcelImport(unittest.TestCase):
    def test_import_twice_no_duplicate(self):
        tmp, conn, repo, std = make_env()
        try:
            f = tmp / "第一场.xlsx"
            make_xlsx(f, [[1, "2000", "BA4XXX", "南京栖霞", "K6", "771", "5W"],
                          [2, "2010", "BG4TKI", "安徽芜湖", "705", "原", "10W"]])
            prov = ExcelImportProvider(repo)
            r1 = prov.import_file(f, std)
            self.assertEqual(r1[0], 2)
            total1 = conn.execute("SELECT COUNT(*) c FROM checkins").fetchone()["c"]
            r2 = prov.import_file(f, std)
            total2 = conn.execute("SELECT COUNT(*) c FROM checkins").fetchone()["c"]
            self.assertEqual(total1, total2, "重复导入不得产生重复记录")
        finally:
            conn.close()

    def test_import_standardizes(self):
        tmp, conn, repo, std = make_env()
        try:
            f = tmp / "b.xlsx"
            make_xlsx(f, [[1, "2000", "BA4XXX", "江苏省南京市栖霞区", "八重洲 FT-1907R", "2.2米玻璃钢", "5W"]])
            ExcelImportProvider(repo).import_file(f, std)
            c = conn.execute("SELECT * FROM checkins").fetchone()
            self.assertEqual(c["qth_standard"], "南京栖霞")
            self.assertEqual(c["device_standard"], "YAESU FT-1907R")
            self.assertEqual(c["antenna_standard"], "2.2米玻璃钢")
            self.assertEqual(c["power_standard"], "5W")
        finally:
            conn.close()


class TestCommitUndo(unittest.TestCase):
    def test_commit_and_soft_delete(self):
        tmp, conn, repo, std = make_env()
        try:
            s = repo.create_session("测试", "2026-08-08")
            c = Checkin(session_id=s.id, sequence_no=1, callsign="BG4TKI",
                        qth_standard="南京栖霞", source="local")
            repo.add_checkin(c)
            repo.update_profiles_from_checkin(c)
            self.assertEqual(repo.next_sequence(s.id), 2)
            repo.soft_delete_checkin(c.id)
            self.assertEqual(len(repo.list_checkins(s.id)), 0)
            # 画像不应因软删除而丢失（历史统计保留）
            self.assertGreaterEqual(len(repo.profiles_for("BG4TKI", "qth")), 1)
        finally:
            conn.close()

    def test_export_rebuild(self):
        tmp, conn, repo, std = make_env()
        try:
            s = repo.create_session("重建", "2026-08-08", operator_callsign="BA4THG",
                                    repeater_name="江苏省中继")
            repo.add_checkin(Checkin(session_id=s.id, sequence_no=1,
                                     checkin_time="2026-08-08T22:13:00",
                                     callsign="BG4TKI", qth_standard="南京栖霞",
                                     device_standard="K6", power_standard="5W",
                                     source="local"))
            dest = tmp / "rebuild.xlsx"
            export_session(repo, s.id, str(dest))
            wb = load_workbook(dest)
            ws = wb.active
            rows = list(ws.iter_rows(values_only=True))
            # 头部：中继/场次/日期、主控、空行，然后才是表头
            self.assertIn("主控", rows[1][0])
            self.assertIn("BA4THG", rows[1][0])
            self.assertEqual(rows[3][2], "呼号")
            self.assertEqual(rows[4][2], "BG4TKI")
            self.assertEqual(rows[4][1], "22:13")  # 时间带冒号
        finally:
            conn.close()


class _FakeProvider:
    name = "fake"

    def __init__(self, mode: str):
        self.mode = mode
        self.uid = "TESTUID"

    def fetch_ranking(self):
        if self.mode == "network_error":
            import requests
            raise requests.ConnectionError("dns failed")
        return {"stats": 1, "list": [{"name": "BA4XXX", "count": 1}]}

    def fetch_many_stats(self, callsigns, progress=None):
        return []

    def select_callsigns(self, ranking_list, local, first_sync):
        return []


class TestSyncTolerance(unittest.TestCase):
    def test_network_error_not_blocking(self):
        tmp, conn, repo, std = make_env()
        try:
            svc = SyncService(repo, _FakeProvider("network_error"), std)
            result = svc.sync_once()
            self.assertFalse(result["ok"])
            self.assertIn("365dt", result["message"])
            state = repo.get_sync_state("365dt")
            self.assertEqual(state.status, "error")
            # 数据库仍可用
            repo.create_session("仍可用", "2026-08-08")
        finally:
            conn.close()

    def test_empty_sync_ok(self):
        tmp, conn, repo, std = make_env()
        try:
            svc = SyncService(repo, _FakeProvider("ok"), std)
            result = svc.sync_once()
            self.assertTrue(result["ok"])
        finally:
            conn.close()


class TestAutoSession(unittest.TestCase):
    def test_commit_auto_creates_session(self):
        """无当前场次时回车提交应自动建场并写入（修复“键入没记录”）。"""
        from tests.helpers import make_service

        svc = make_service()
        try:
            self.assertIsNone(svc.current_session())
            r = svc.parse("bg4tki njqx k6 y 5")
            res = svc.commit(r)
            self.assertTrue(res["ok"])
            self.assertIsNotNone(svc.current_session())
            self.assertEqual(res["checkin"].callsign, "BG4TKI")
            self.assertEqual(len(svc.list_checkins()), 1)
        finally:
            svc.close()


def _mk_checkin(seq: int, callsign: str, qth: str, time: str = "2026-08-08T22:13:00") -> Checkin:
    return Checkin(sequence_no=seq, callsign=callsign, qth_standard=qth,
                   device_standard="K6", antenna_standard="771", power_standard="5W",
                   checkin_time=time)


class TestConsistency(unittest.TestCase):
    """SQLite/Excel 一致性比对（P0-6）。"""

    def test_pass(self):
        from services.app_service import build_consistency_report
        sqlite = [_mk_checkin(1, "BG4TKI", "南京栖霞"), _mk_checkin(2, "BA4XXX", "安徽芜湖")]
        excel = [{"sequence": "1", "time": "22:13", "callsign": "BG4TKI", "qth": "南京栖霞",
                  "device": "K6", "antenna": "771", "power": "5W"},
                 {"sequence": "2", "time": "22:13", "callsign": "BA4XXX", "qth": "安徽芜湖",
                  "device": "K6", "antenna": "771", "power": "5W"}]
        ok, report = build_consistency_report(sqlite, excel)
        self.assertTrue(ok)
        self.assertIn("PASS", report)

    def test_diff(self):
        from services.app_service import build_consistency_report
        sqlite = [_mk_checkin(1, "BG4TKI", "南京栖霞"),
                  _mk_checkin(2, "BA4XXX", "安徽芜湖"),
                  _mk_checkin(3, "BD4ABC", "南京鼓楼")]
        excel = [{"sequence": "1", "time": "2213", "callsign": "BG4TKI", "qth": "南京鼓楼",
                  "device": "K6", "antenna": "771", "power": "5W"},
                 {"sequence": "3", "time": "2213", "callsign": "BD4ABC", "qth": "南京鼓楼",
                  "device": "K6", "antenna": "771", "power": "5W"}]
        ok, report = build_consistency_report(sqlite, excel)
        self.assertFalse(ok)
        self.assertIn("#2", report)  # Excel 缺失
        self.assertIn("#1", report)  # QTH 不同


class TestMigrationV2(unittest.TestCase):
    def test_v2_columns_and_unsynced_tracking(self):
        tmp, conn, repo, std = make_env()
        try:
            cols = [r["name"] for r in conn.execute("PRAGMA table_info(checkins)")]
            self.assertIn("excel_synced", cols)
            self.assertIn("excel_row", cols)
            s = repo.create_session("t", "2026-08-08")
            c = Checkin(session_id=s.id, sequence_no=1, callsign="BG4TKI", source="local")
            repo.add_checkin(c)
            self.assertIn(c.id, [x.id for x in repo.list_unsynced(s.id)])
            repo.mark_excel_synced(c.id, 5)
            self.assertEqual(len(repo.list_unsynced(s.id)), 0)
        finally:
            conn.close()


def make_service():
    from tests.helpers import make_service as _make_service

    return _make_service()


class TestV2Features(unittest.TestCase):
    def test_complete_prefix(self):
        svc = make_service()
        try:
            svc.create_session("t", "2026-08-08")
            hits = [v for _, v in svc.complete("nj")]
            self.assertTrue(any("南京" in h for h in hits), hits[:3])
            labels = [k for k, _ in svc.complete("id")]
            self.assertTrue(any("ICOM" in l for l in labels), labels[:3])
            # 补全插入的应是可再次解析的 key（不是完整标准名）
            values = [v for _, v in svc.complete("id")]
            self.assertIn("id52", values)
        finally:
            svc.close()

    def test_suggest_qth_abbr(self):
        svc = make_service()
        try:
            abbr, conflicts = svc.suggest_qth_abbr("南京栖霞")
            self.assertEqual(abbr, "njqx")
            self.assertEqual(conflicts, [])
            abbr2, _ = svc.suggest_qth_abbr("浙江杭州西湖")
            self.assertEqual(abbr2, "zjhzxh")
            # 制造冲突：njqx 已被占为别的 QTH
            svc.set_alias("qth", "njqx", "南京鼓楼")
            _, conflicts2 = svc.suggest_qth_abbr("南京栖霞")
            self.assertIn("南京鼓楼", conflicts2)
        finally:
            svc.close()

    def test_update_checkin_syncs(self):
        svc = make_service()
        try:
            sess = svc.create_session("t", "2026-08-08")
            r = svc.parse("bg4tki njqx k6 y 5")
            res = svc.commit(r)
            cid = res["checkin"].id
            out = svc.update_checkin(cid, "qth", "南京鼓楼")
            self.assertTrue(out["ok"])
            c = svc.repo.get_checkin(cid)
            self.assertEqual(c.qth_standard, "南京鼓楼")
            audit = svc.repo.list_audit(cid)
            self.assertTrue(any(a.field_name == "qth" for a in audit))
            # 画像已重算
            prof = [p.field_value for p in svc.repo.profiles_for("BG4TKI", "qth")]
            self.assertEqual(prof, ["南京鼓楼"])
        finally:
            svc.close()

    def test_suggest_aliases_from_imports(self):
        """导入历史 → 自动生成词典建议（人工确认）。"""
        svc = make_service()
        try:
            s = svc.create_session("t", "2026-08-08")
            for _ in range(2):
                svc.repo.add_checkin(Checkin(
                    session_id=s.id, sequence_no=svc.repo.next_sequence(s.id),
                    checkin_time="2026-08-08T10:00:00", callsign="BA4XXX",
                    device_raw="KENWOOD TM-V71A", device_standard="KENWOOD TM-V71A",
                    qth_standard="盐城", source="365dt"))
            sugg = svc.suggest_aliases_from_imports()
            dev = [x for x in sugg if x["kind"] == "device"]
            self.assertTrue(any(x["alias"] == "tmv71a" and x["standard"] == "KENWOOD TM-V71A"
                                for x in dev), dev[:3])
            qth = [x for x in sugg if x["kind"] == "qth"]
            self.assertTrue(any(x["alias"] == "yc" and x["standard"] == "盐城" for x in qth), qth[:3])
            # 加入后即可解析
            svc.add_alias_suggestions([x for x in sugg if x["alias"] == "tmv71a"])
            r = svc.parse("ba4xxx tmv71a")
            self.assertEqual(r.device.value, "KENWOOD TM-V71A")
        finally:
            svc.close()


class TestPerformance(unittest.TestCase):
    def test_150_commits(self):
        """连续 150 条录入无明显性能下降（V1 实测项）。"""
        import time as _time
        svc = make_service()
        try:
            svc.create_session("压测", "2026-08-08")
            t0 = _time.time()
            for i in range(150):
                callsign = f"BG{i:03d}" if i % 2 else f"BA{i:03d}"
                r = svc.parse(f"{callsign.lower()} njqx k6 y 5")
                res = svc.commit(r)
                self.assertTrue(res["ok"])
            elapsed = _time.time() - t0
            self.assertEqual(svc.repo.next_sequence(svc.current_session().id), 151)
            # 性能冒烟：负载波动时放宽，防止 flaky（正确性由上面断言保证）
            self.assertLess(elapsed, 30.0, f"150 条耗时 {elapsed:.2f}s 偏慢")
        finally:
            svc.close()


class TestTestIsolation(unittest.TestCase):
    """任务书第一阶段 #1：测试/benchmark 不得修改真实 config.json。"""

    def test_tests_do_not_modify_real_config(self):
        """运行全套测试前后，真实 config.json SHA256 必须一致。"""
        from tests.helpers import make_service, real_config_sha

        before = real_config_sha()
        svc = make_service()
        try:
            s = svc.create_session("隔离", "2026-08-08")
            for i in range(3):
                r = svc.parse(f"bg4tki{i} njqx k6 y 5")
                svc.commit(r)
            svc.set_alias("qth", "zztest", "郑州")
            svc.settings.set("default_province", "安徽")
            svc.end_current_session()
        finally:
            svc.close()
        from tests.helpers import assert_real_config_untouched

        assert_real_config_untouched(self, before)

    def test_temp_settings_do_not_touch_app_config(self):
        """用临时路径创建的 Settings，写值/保存后真实 config.json 不变。"""
        import hashlib

        from config.settings import CONFIG_PATH
        from tests.helpers import make_settings

        before = hashlib.sha256(CONFIG_PATH.read_bytes()).hexdigest() if CONFIG_PATH.exists() else None
        s, tmp = make_settings()
        s.set_many(default_province="安徽", fuzzy_high=95.0)
        s.set("dt365_uid", "TEST_UID")
        s.save()
        # 确认配置写到临时目录而非真实位置
        self.assertTrue((tmp / "config.json").exists())
        after = hashlib.sha256(CONFIG_PATH.read_bytes()).hexdigest() if CONFIG_PATH.exists() else None
        self.assertEqual(before, after, "临时 Settings 不应改动真实 config.json")

    def test_benchmark_uses_isolated_config(self):
        """bench 核心逻辑必须使用隔离 Settings（临时 config 路径）。"""
        from tests.helpers import make_settings, real_config_sha

        before = real_config_sha()
        s, tmp = make_settings()
        # 模拟 bench.py 的用法：隔离 config + 隔离数据目录
        s.set_many(data_dir=str(tmp / "data"), logs_dir=str(tmp / "logs"),
                   backup_dir=str(tmp / "backup"))
        from services.app_service import AppService

        svc = AppService(s)
        try:
            svc.create_session("bench", "2026-08-08")
            for i in range(10):
                svc.commit(svc.parse(f"bg{i:02d}tki njqx k6 y 5"))
        finally:
            svc.close()
        from tests.helpers import assert_real_config_untouched

        assert_real_config_untouched(self, before)


if __name__ == "__main__":
    unittest.main()
