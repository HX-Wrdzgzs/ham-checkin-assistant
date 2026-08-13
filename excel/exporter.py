"""Excel 导出：SQLite → 新 xlsx（即使原点名表丢失也能恢复，规格第 50 节）。"""
from __future__ import annotations

from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font

from database.repository import Repository
from excel.template import EXPORT_HEADERS


def hhmm(iso_time: str) -> str:
    """'2026-08-08T18:16:27+08:00' → '18:16'"""
    if not iso_time:
        return ""
    try:
        dt = datetime.fromisoformat(iso_time)
        return dt.strftime("%H:%M")
    except ValueError:
        return iso_time[:5] if iso_time else ""


def export_session(repo: Repository, session_id: int, dest_path: str) -> str:
    """导出本场签到为 xlsx（含主控/中继/日期头部），返回写入路径。"""
    session = repo.get_session(session_id)
    checkins = repo.list_checkins(session_id)
    wb = Workbook()
    ws = wb.active
    ws.title = session.name or f"场次{session_id}"
    # 头部信息：中继/场次/日期/主控
    head = [
        f"{session.repeater_name or '点名'} · {session.name or ''} · {session.date or ''}",
        f"主控：{session.operator_callsign or '（未设置）'}",
        "",
    ]
    for row in head:
        ws.append([row])
    ws.append(EXPORT_HEADERS)
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)
    for c in checkins:
        ws.append([
            c.sequence_no, hhmm(c.checkin_time), c.callsign,
            c.qth_standard, c.device_standard, c.antenna_standard,
            c.power_standard, c.signal, c.source,
        ])
    # 列宽
    for col, width in zip("ABCDEFGHI", (6, 8, 12, 16, 16, 14, 8, 8, 12)):
        ws.column_dimensions[col].width = width
    wb.save(dest_path)
    return dest_path


def export_template(path: str) -> str:
    """生成空白点名模板。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "点名记录"
    ws.append(EXPORT_HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for col, width in zip("ABCDEFGHI", (6, 8, 12, 16, 16, 14, 8, 8, 12)):
        ws.column_dimensions[col].width = width
    wb.save(path)
    return path
